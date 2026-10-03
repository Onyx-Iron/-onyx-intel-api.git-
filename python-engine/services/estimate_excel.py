"""
Estimate workbook export with live Excel formulas (openpyxl).

Writes =A2*B2 style formulas so GCs can audit markups inside Excel without
flattening calculated values.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font


def build_estimate_workbook(payload: dict[str, Any]) -> bytes:
    """
    payload = {
      project_name: str,
      overhead_pct: float,
      profit_pct: float,
      contingency_pct: float,
      rows: [{cost_code, description, quantity, unit,
              labor_unit, material_unit, equipment_unit,
              subcontractor_unit, trucking_unit, disposal_unit}, ...]
    }
    """
    project = str(payload.get("project_name") or "Estimate")
    overhead = float(payload.get("overhead_pct") or 0) / 100.0
    profit = float(payload.get("profit_pct") or 0) / 100.0
    contingency = float(payload.get("contingency_pct") or 0) / 100.0
    rows: list[dict[str, Any]] = list(payload.get("rows") or [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule of Values"

    headers = [
        "Item #",
        "Cost Code",
        "Description",
        "Qty",
        "Unit",
        "Labor Unit",
        "Material Unit",
        "Equipment Unit",
        "Sub Unit",
        "Trucking Unit",
        "Disposal Unit",
        "Labor",
        "Material",
        "Equipment",
        "Sub",
        "Trucking",
        "Disposal",
        "Direct",
        "Overhead",
        "Profit",
        "SOV Value",
    ]
    ws.append([f"Schedule of Values · {project}"])
    ws.append([])
    ws.append(headers)
    for cell in ws[3]:
        cell.font = Font(bold=True)

    # Data starts at Excel row 4
    first_data = 4
    for i, r in enumerate(rows):
        excel_row = first_data + i
        ws.append([
            i + 1,
            r.get("cost_code") or "",
            r.get("description") or "",
            float(r.get("quantity") or 0),
            r.get("unit") or "",
            float(r.get("labor_unit") or 0),
            float(r.get("material_unit") or 0),
            float(r.get("equipment_unit") or 0),
            float(r.get("subcontractor_unit") or 0),
            float(r.get("trucking_unit") or 0),
            float(r.get("disposal_unit") or 0),
            # Live formulas: Qty * Unit rates
            f"=D{excel_row}*F{excel_row}",  # Labor
            f"=D{excel_row}*G{excel_row}",  # Material
            f"=D{excel_row}*H{excel_row}",  # Equipment
            f"=D{excel_row}*I{excel_row}",  # Sub
            f"=D{excel_row}*J{excel_row}",  # Trucking
            f"=D{excel_row}*K{excel_row}",  # Disposal
            f"=SUM(L{excel_row}:Q{excel_row})",  # Direct
            f"=R{excel_row}*{overhead}",  # Overhead
            f"=(R{excel_row}+S{excel_row})*{profit}",  # Profit
            f"=R{excel_row}+S{excel_row}+T{excel_row}",  # SOV
        ])

    last_data = first_data + len(rows) - 1 if rows else first_data - 1
    total_row = last_data + 2 if rows else first_data + 1
    ws.cell(total_row, 3, "TOTAL")
    if rows:
        for col, letter in enumerate(["L", "M", "N", "O", "P", "Q", "R", "S", "T", "U"], start=12):
            ws.cell(total_row, col, f"=SUM({letter}{first_data}:{letter}{last_data})")

    # Proposal sheet with live rollups
    prop = wb.create_sheet("Proposal")
    prop.append([f"Proposal · {project}"])
    prop.append([])
    prop.append(["Metric", "Value"])
    prop["A3"].font = Font(bold=True)
    prop["B3"].font = Font(bold=True)
    if rows:
        prop.append(["Direct Cost Total", f"='Schedule of Values'!R{total_row}"])
        prop.append([f"Contingency ({contingency * 100:.1f}%)", f"=B4*{contingency}"])
        prop.append(["Subtotal", "=B4+B5"])
        prop.append([f"Overhead ({overhead * 100:.1f}%)", f"=B6*{overhead}"])
        prop.append([f"Profit ({profit * 100:.1f}%)", f"=(B6+B7)*{profit}"])
        prop.append(["FINAL BID", "=B6+B7+B8"])
        prop["A9"].font = Font(bold=True)
        prop["B9"].font = Font(bold=True)
    else:
        prop.append(["Direct Cost Total", 0])
        prop.append(["FINAL BID", 0])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
